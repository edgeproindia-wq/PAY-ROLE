"""Owner page: bulk-import Employee Master from Excel (.xlsx) or CSV.
Step 1 upload -> every row is validated with the same rules as the Add Employee
form and shown as a preview (nothing saved). Step 2 "Import" saves the valid rows
in one transaction, optionally creating Employee ID logins with one password."""
import csv
import datetime
import io
import re

from django import forms
from django.contrib import messages
from django.contrib.auth.password_validation import validate_password
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render

from .audit import log_action
from .forms import EmployeeForm
from .models import Employee, User
from .permissions import company_owner_required

# (field, header shown in the template, required, accepted header spellings)
COLUMNS = [
    ('employee_code', 'Employee ID', True, ['employee id', 'emp code', 'employee code', 'emp id', 'code', 'empcode']),
    ('first_name', 'First Name', True, ['first name', 'firstname', 'name', 'employee name', 'emp name', 'full name']),
    ('last_name', 'Last Name', False, ['last name', 'lastname', 'surname']),
    ('email', 'Email', True, ['email', 'email id', 'e-mail', 'mail id', 'official email']),
    ('phone', 'Phone', False, ['phone', 'mobile', 'mobile no', 'phone no', 'contact', 'contact no']),
    ('gender', 'Gender (M/F/O)', False, ['gender', 'gender (m/f/o)', 'sex']),
    ('date_of_birth', 'Date of Birth', False, ['date of birth', 'dob', 'birth date']),
    ('date_of_joining', 'Date of Joining', True, ['date of joining', 'doj', 'joining date', 'join date']),
    ('department', 'Department', True, ['department', 'dept']),
    ('designation', 'Designation', True, ['designation', 'role', 'title', 'position']),
    ('employment_status', 'Status', False, ['status', 'employment status']),
    ('pan_number', 'PAN', False, ['pan', 'pan no', 'pan number']),
    ('aadhar_number', 'Aadhaar', False, ['aadhaar', 'aadhar', 'aadhaar no', 'aadhar no', 'aadhaar number', 'aadhar number']),
    ('bank_name', 'Bank Name', False, ['bank name', 'bank']),
    ('account_holder_name', 'Account Holder Name', False, ['account holder name', 'account holder', 'holder name']),
    ('bank_account_no', 'Bank Account No', False, ['bank account no', 'account no', 'account number', 'bank account number', 'a/c no']),
    ('ifsc_code', 'IFSC', False, ['ifsc', 'ifsc code']),
    ('bank_branch', 'Bank Branch', False, ['bank branch', 'branch']),
    ('account_type', 'Account Type (SAVINGS/CURRENT/SALARY)', False, ['account type', 'account type (savings/current/salary)']),
]
SESSION_KEY = 'employee_import_rows'
MAX_ROWS = 2000
GENDER = {'m': 'M', 'male': 'M', 'f': 'F', 'female': 'F', 'o': 'O', 'other': 'O'}
STATUS = {'active': 'ACTIVE', 'on leave': 'ON_LEAVE', 'on_leave': 'ON_LEAVE', 'resigned': 'RESIGNED',
          'terminated': 'TERMINATED', 'inactive': 'RESIGNED'}


def _norm(h):
    return re.sub(r'\s+', ' ', str(h or '').strip().lower().replace('_', ' '))


def _date(v):
    if v in (None, ''):
        return ''
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d.%m.%Y', '%d-%b-%Y', '%d %b %Y', '%d-%b-%y', '%m/%d/%Y'):
        try:
            return datetime.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return s                                   # the form will report it as invalid


def _cell(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)                              # Excel numbers: 9876543210.0 -> 9876543210
    return str(v).strip()


def read_rows(upload):
    """Return (rows as dicts keyed by model field, error message)."""
    name = upload.name.lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith('.csv'):
        text = upload.read().decode('utf-8-sig', errors='replace')
        raw = list(csv.reader(io.StringIO(text)))
    else:
        return [], 'Upload an Excel (.xlsx) or CSV file.'
    raw = [r for r in raw if any(_cell(c) for c in r)]
    if not raw:
        return [], 'The file is empty.'
    header = [_norm(h) for h in raw[0]]
    index = {}
    for field, label, _req, aliases in COLUMNS:
        for i, h in enumerate(header):
            if h == _norm(label) or h in aliases:
                index[field] = i
                break
    missing = [label for field, label, req, _ in COLUMNS if req and field not in index]
    if missing:
        return [], 'Missing column(s): ' + ', '.join(missing) + '. Download the template to see the expected columns.'
    if len(raw) - 1 > MAX_ROWS:
        return [], f'Too many rows (max {MAX_ROWS}).'
    rows = []
    for r in raw[1:]:
        d = {}
        for field, i in index.items():
            v = r[i] if i < len(r) else ''
            d[field] = _date(v) if field in ('date_of_birth', 'date_of_joining') else _cell(v)
        if d.get('first_name') and not index.get('last_name') and ' ' in d['first_name']:
            d['first_name'], d['last_name'] = d['first_name'].split(' ', 1)   # "Full Name" column
        d['gender'] = GENDER.get(d.get('gender', '').lower(), d.get('gender', '').upper())
        d['employment_status'] = STATUS.get(d.get('employment_status', '').lower(), d.get('employment_status', '').upper() or 'ACTIVE')
        d['account_type'] = (d.get('account_type') or '').upper()
        rows.append(d)
    return rows, ''


def validate_rows(rows, company):
    """Same rules as Add Employee + unique Employee ID per company + unique email."""
    report, seen_codes, seen_emails = [], set(), set()
    for n, d in enumerate(rows, start=2):
        data = {f: d.get(f, '') for f, *_ in COLUMNS}
        data['bank_status'] = 'UNVERIFIED'            # new bank details always start unverified
        data['employment_status'] = data.get('employment_status') or 'ACTIVE'
        form = EmployeeForm(data=data)
        errors = []
        if not form.is_valid():
            for field, errs in form.errors.items():
                label = next((c[1] for c in COLUMNS if c[0] == field), field)
                errors.append(f'{label}: {" ".join(errs)}')
        code = (d.get('employee_code') or '').strip()
        email = (d.get('email') or '').strip().lower()
        if code:
            if code.lower() in seen_codes:
                errors.append('Employee ID repeated in the file')
            elif Employee.objects.filter(company=company, employee_code__iexact=code).exists():
                errors.append('Employee ID already exists in Employee Master')
            seen_codes.add(code.lower())
        if email:
            if email in seen_emails:
                errors.append('Email repeated in the file')
            seen_emails.add(email)
        report.append({'line': n, 'code': code, 'name': f"{d.get('first_name', '')} {d.get('last_name', '')}".strip(),
                       'errors': errors, 'data': data})
    return report


class UploadForm(forms.Form):
    file = forms.FileField(help_text='Excel (.xlsx) or CSV, first row = column headings')


class ConfirmForm(forms.Form):
    create_logins = forms.BooleanField(required=False, initial=True, label='Also create Employee ID logins')
    password1 = forms.CharField(widget=forms.PasswordInput, required=False, label='Login password for everyone')
    password2 = forms.CharField(widget=forms.PasswordInput, required=False, label='Confirm password')

    def clean(self):
        c = super().clean()
        if c.get('create_logins'):
            p1, p2 = c.get('password1'), c.get('password2')
            if not p1:
                raise forms.ValidationError('Enter the login password, or untick "Also create Employee ID logins".')
            if p1 != p2:
                raise forms.ValidationError('The two passwords do not match.')
            try:
                validate_password(p1)
            except forms.ValidationError as e:
                raise forms.ValidationError(e.messages)
        return c


def _username(e):
    base = f"emp{e.company_id}-{e.employee_code}".lower().replace(' ', '')[:140]
    username, n = base, 1
    while User.objects.filter(username__iexact=username).exists():
        n += 1
        username = f"{base}-{n}"
    return username


@company_owner_required
def hr_employee_import(request):
    company = request.user.company
    upload = UploadForm()
    confirm = ConfirmForm(initial={'create_logins': True})
    report = None
    if request.method == 'POST' and request.POST.get('step') == 'preview':
        upload = UploadForm(request.POST, request.FILES)
        if upload.is_valid():
            rows, err = read_rows(upload.cleaned_data['file'])
            if err:
                upload.add_error('file', err)
            else:
                report = validate_rows(rows, company)
                request.session[SESSION_KEY] = [r['data'] for r in report if not r['errors']]
    elif request.method == 'POST' and request.POST.get('step') == 'import':
        confirm = ConfirmForm(request.POST)
        rows = request.session.get(SESSION_KEY) or []
        if not rows:
            messages.error(request, 'Nothing to import - upload the file again.')
            return redirect('hr_employee_import')
        if confirm.is_valid():
            fresh = validate_rows(rows, company)           # re-check: data may have changed since the preview
            good = [r['data'] for r in fresh if not r['errors']]
            created = logins = 0
            with transaction.atomic():
                for data in good:
                    form = EmployeeForm(data=data)
                    form.is_valid()
                    e = form.save(commit=False)
                    e.company = company
                    e.save()
                    created += 1
                    if confirm.cleaned_data['create_logins']:
                        e.user = User.objects.create_user(
                            username=_username(e), password=confirm.cleaned_data['password1'], role='EMPLOYEE',
                            company=company, email=e.email, first_name=e.first_name, last_name=e.last_name)
                        e.save(update_fields=['user'])
                        logins += 1
            request.session.pop(SESSION_KEY, None)
            log_action(request, 'CREATE', company, details=f'Employee import: {created} employees, {logins} logins')
            skipped = len(rows) - len(good)
            messages.success(request, f'{created} employee(s) imported' + (f', {logins} login(s) created' if logins else '')
                             + (f'. {skipped} row(s) skipped because they are no longer valid.' if skipped else '.'))
            return redirect('employee_master')
        report = validate_rows(rows, company)
    for f in list(upload.fields.values()) + list(confirm.fields.values()):
        if not isinstance(f.widget, forms.CheckboxInput):
            f.widget.attrs.setdefault('class', 'form-control')
    ok = [r for r in report if not r['errors']] if report is not None else []
    bad = [r for r in report if r['errors']] if report is not None else []
    return render(request, 'features/hr_employee_import.html',
                  {'upload': upload, 'confirm': confirm, 'report': report, 'ok': ok, 'bad': bad})


@company_owner_required
def hr_employee_import_template(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Employees'
    ws.append([c[1] for c in COLUMNS])
    ws.append(['EPRO0001', 'Anu', 'Kumar', 'anu.kumar@company.com', '9876543210', 'F', '1995-04-12', '2024-06-03',
               'Detailing', 'Tekla Modeler', 'ACTIVE', 'ABCDE1234F', '123412341234', 'HDFC Bank', 'Anu Kumar',
               '123456789012', 'HDFC0001234', 'Hosur', 'SAVINGS'])
    for i, c in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4A4AD8' if c[2] else '64748B')
        ws.column_dimensions[cell.column_letter].width = max(14, len(c[1]) + 2)
    notes = wb.create_sheet('Instructions')
    for line in ['Fill one employee per row on the Employees sheet (delete the example row).',
                 'Purple headings are required: Employee ID, First Name, Email, Date of Joining, Department, Designation.',
                 'Dates: 2024-06-03 or 03-06-2024 or 03/06/2024.',
                 'Gender: M, F or O. Status: ACTIVE, ON_LEAVE, RESIGNED or TERMINATED (blank = ACTIVE).',
                 'Each Email must be unique. IFSC like HDFC0001234; account number 9-18 digits; Aadhaar 12 digits.',
                 'Employees whose Employee ID already exists are skipped (not changed).']:
        notes.append([line])
    notes.column_dimensions['A'].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="employee_import_template.xlsx"'
    return resp